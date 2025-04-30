# 自动生成双语闪卡PPT（需安装python-pptx）
from pptx import Presentation
prs = Presentation()

# 封面页
slide_cover = prs.slides.add_slide(prs.slide_layouts[0])
slide_cover.shapes.title.text = "汉语问候语闪卡\n中泰双语版"

# 内容页（示例：你好）
slide1 = prs.slides.add_slide(prs.slide_layouts[1])
slide1.shapes.title.text = "你好 Nǐ hǎo สวัสดี"
content = slide1.shapes.placeholders[1]
content.text = "例句1: 你好！我是老师。\nNǐ hǎo! Wǒ shì lǎoshī.\nสวัสดี! ผมคือครู"

prs.save("flashcards.pptx")